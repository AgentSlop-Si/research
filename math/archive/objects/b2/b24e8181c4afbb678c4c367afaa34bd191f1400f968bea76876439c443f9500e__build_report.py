from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

P=Path(__file__).resolve().parent
text=(P/'REPORT.txt').read_text()
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
normal=doc.styles['Normal']
normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(4)
normal.paragraph_format.line_spacing=1.02
for name in ('Title','Heading 1'):
    doc.styles[name].font.color.rgb=RGBColor(0,0,0)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Title'].font.size=Pt(20)
doc.styles['Heading 1'].font.size=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
parts=text.split('\n\n')
first=parts.pop(0).splitlines()
doc.add_paragraph('Complex Fourier network label compression',style='Title')
doc.add_paragraph(first[1])
src=0
for part in parts:
    lines=part.splitlines()
    if not lines:continue
    if len(lines)==1 and part.isupper():
        doc.add_paragraph(part.capitalize(),style='Heading 1')
        if part=='PRIMARY SOURCES':break
        continue
    if any(line.startswith('https://') for line in lines):
        for url in lines:
            if not url.startswith('https://'):
                doc.add_paragraph(url);continue
            src+=1;p=doc.add_paragraph()
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'12608F');rp.append(color);r.append(rp)
            t=OxmlElement('w:t');t.text='Open pinned source '+str(src);r.append(t);h.append(r);p._p.append(h)
    elif any(line.startswith('  ') for line in lines):
        p=doc.add_paragraph('\n'.join(lines))
        p.paragraph_format.left_indent=Inches(.12)
        for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
    else:doc.add_paragraph(' '.join(lines))
sources=[line for line in text.splitlines() if line.startswith('https://')]
p=doc.add_paragraph()
for i,(label,url) in enumerate(zip(['Producer and labels','Baseline histogram','Copied centre interface','Attribution'],sources)):
    if i:p.add_run('  |  ')
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'12608F');rp.append(color);r.append(rp)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
for p in doc.paragraphs:
    for border in list(p._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='Complex Fourier network label compression'
doc.core_properties.subject='Bounded exact label and moment certificate audit'
doc.core_properties.author=''
doc.save(P/'Complex_Fourier_label_audit.docx')

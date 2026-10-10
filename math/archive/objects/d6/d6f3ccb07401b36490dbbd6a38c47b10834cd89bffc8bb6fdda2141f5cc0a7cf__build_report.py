from pathlib import Path
import json, math, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
py=json.loads((ROOT/'python_experiment/python_results.json').read_text())
first=json.loads((ROOT/'c_experiment/results.json').read_text())
confirm=ROOT/'c_experiment/results_confirmation.json'
c=json.loads(confirm.read_text()) if confirm.exists() else first
assert len(c['sizes'])==5, 'Wait for complete native run'
out=ROOT/'deliverables'; out.mkdir(exist_ok=True)
d=Document(); sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt(10.5 if name=='Normal' else {'Title':23,'Subtitle':11,'Heading 1':15,'Heading 2':12}[name])
 s.paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.08

def p(text): return d.add_paragraph(text)
def h(text): d.add_heading(text,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths): t.columns[i].width=Inches(w)
 for i,head in enumerate(headers):t.rows[0].cells[i].text=head
 for row in rows:
  cells=t.add_row().cells
  for cell,txt in zip(cells,row):cell.text=str(txt)
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=cell._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ('top','left','bottom','right'):
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'D9E5F2' if ri==0 else ('F7F9FB' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
   mar=OxmlElement('w:tcMar')
   for side in ('top','left','bottom','right'):
    item=OxmlElement('w:'+side);item.set(qn('w:w'),'85');item.set(qn('w:type'),'dxa');mar.append(item)
   tcPr.append(mar)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.05
    pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.RIGHT
    for rr in pp.runs:rr.font.size=Pt(9);rr.bold=ri==0
  trPr=row._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');trPr.append(keep)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 p('')
 return t

adaptive_file=ROOT/'c_experiment/adaptive_results.json'
if adaptive_file.exists():
 a=json.loads(adaptive_file.read_text())
 d.add_paragraph('Adaptive binary FFT implementation','Title')
 p('Runnable implementation and held out validation | 8 October 2026')
 p('The implementation now offers a reusable binary FFT plan with a size dependent kernel choice. For one or a few transforms, use the explicit fixed arithmetic plan to avoid calibration. Optional autotuning compares arithmetic and small lookup kernels once, then reuses its decision. It chose the fastest held-out kernel at three of five sizes; its conservative rule skipped smaller gains at the other two.')
 h('Use the fast path first')
 p('For one-off or short workloads, construct BinaryFFTPlan(n, method="codelet8") and call plan.execute(x), optionally with out= for a caller-owned output buffer. This path performs no calibration. Inputs must be contiguous one-dimensional NumPy uint8 or bool arrays containing only zero and one. Outputs are complex128; new outputs are independent by default.')
 p('For heavy repeated use, get_plan(n, dtype) or binary_fft(x) opts into cached autotuning. The first call is expensive. Up to eight plans are cached in process by size and dtype. Repeated execution does not retune; plan.retune() is explicit. The README contains complete runnable examples and close/clear-cache behavior.')
 h('What the selector chose')
 rows=[]
 for case in a['sizes']:
  name=lambda v:v.replace('codelet','arithmetic ').replace('lookup','lookup ')
  rows.append([f"{case['n']:,}",name(case['method']),name(case['heldout_fastest']),f"{case['candidate_choice_regret_percent']:.1f}%",f"{case['calibration']['elapsed_ms']:.1f}"])
 table(['N','Selected kernel','Held out fastest','Choice penalty','Calibration ms'],rows,[.85,1.35,1.35,1.15,1.05])
 p('Choice penalty compares the same candidate kernels on held-out timing samples, independently of noisy selected-call timing. Autotuning requires at least an 8% lookup advantage and lookup wins in at least 75% of calibration rounds; otherwise it chooses arithmetic. This deliberately rejects marginal wins, but it cannot guarantee the fastest method on future inputs or under changing machine load.')
 h('Calibration is worth paying only with enough reuse')
 p('Measured default calibration cost was roughly 0.44 to 0.81 seconds through N = 262,144, and 4.02 seconds at N = 1,048,576. At N = 4,096 the observed saving versus fixed arithmetic8 took about 317,000 transforms to repay calibration. At N = 65,536 the selector chose arithmetic8, so there is no demonstrated payback over choosing that kernel directly. Use fixed mode unless the expected repeated workload justifies tuning.')
 p('Held-out tests used eight new seeds per size and 13 randomized timing rounds, separate from calibration. All results matched NumPy; worst absolute error was 2.28e-12. Repeated calls kept calibration_count at one, cached plans were reused, and invalid input, output alignment, aliasing, closed-plan and buffer-ownership checks passed. Public API timing includes validation, ctypes dispatch and optional output allocation; raw sample overlap prevents a reliable isolated dispatch-overhead claim.')
 d.add_page_break()

p0=d.add_paragraph('Selective FFT caching benchmark','Title')
p('Measured on 8 October 2026 | Random binary inputs | Single cloud CPU')
p('A fixed lookup table for small binary FFT leaves reduces bookkeeping and clearly improves the Python prototype. In compiled C, the benefit is size dependent and shrinks against a good arithmetic leaf kernel. At the original size of 65,536, an 8-bit table does not establish a dependable improvement over the best arithmetic implementation. A 16-bit table is consistently unattractive at larger sizes.')
h('Compiled results')
p('Median microseconds per full transform. The longer confirmation run is shown below; the first run is retained in the source bundle. All C methods use identical upper FFT stages, preallocated output, and reusable plans.')
rows=[]
for case in c['sizes']:
 m=case['methods']; best=case['best_arithmetic']; t=lambda name:m[name]['median_us']
 rows.append([f"{case['n']:,}",f"{t('radix2_fused2'):,.2f}",f"{t(best):,.2f}",f"{t('lookup4'):,.2f}",f"{t('lookup8'):,.2f}",f"{t('lookup16'):,.2f}",f"{t(best)/t('lookup8'):.2f}x"])
table(['N','Radix 2','Best arithmetic','4 bit lookup','8 bit lookup','16 bit lookup','8 bit speedup'],rows,[.84,.85,1.04,.93,.93,.98,.90])
p('At N = 65,536 the 4-bit table was about 6% faster than the best codelet in the confirmation run, but slightly slower in the first run. That small advantage was not repeatable. A speedup below 1.00x means the lookup version is slower. The arithmetic controls include fused 4, 8 and 16 point real-input codelets; this avoids attributing a win merely to removing recursion or combining small stages.')
h('What the timing includes')
p('The C timing includes input permutation and gathers, packing each binary key, copying its full output row, and every remaining FFT merge. It excludes input generation, output allocation and reusable table or plan creation. The original input is never assumed to arrive already bit-packed or already permuted.')
p('Each native sample cycles through 16 distinct seeded inputs. Methods run in randomized order on one pinned CPU. The initial experiment used 11 samples of about 40 ms; confirmation used 21 samples of about 100 ms. Shared virtualized hardware still introduces noise, so differences of a few percent should not determine a production choice without a target-machine test.')

# Page 2

d.add_page_break();h('Python results and benchmark corrections')
p('Median milliseconds, 15 measurements per method across three seeds. Warm variants reuse the same twiddle plan; lookup variants reuse their table. Garbage collection occurs before every sample and is disabled during that sample to avoid accumulating recursive closure cycles.')
rows=[]
for case in py['cases']:
 m=case['methods']; t=lambda k:m[k]['median_seconds']*1000
 rows.append([f"{case['n']:,}",f"{t('standard_warm'):.3f}",f"{t('full_cache_no_stats_warm'):.3f}",f"{t('lookup8_warm'):.3f}",f"{t('lookup16_warm'):.3f}",f"{t('standard_warm')/t('lookup8_warm'):.2f}x"])
table(['N','Uncached','Full cache stripped','8 bit lookup','16 bit lookup','8 bit speedup'],rows,[.82,1.03,1.1,1.04,1.08,.95])
p('The original cached prototype increments a counter for every butterfly, counts every tree visit, and builds summary statistics; its uncached comparison does not. The stripped variant removes those counters and summaries while retaining the original dictionaries and DAG. Original unchanged timings are included in the raw data. Full caching moved around break-even in this environment, rather than consistently reproducing the earlier reported slowdown.')
h('Storage and one time setup')
rows=[]
for b in (4,8,16):
 t=c['tables'][str(b)];pt=py['tables'][str(b)]
 rows.append([b,f"{1<<b:,}", {4:'1 KiB',8:'32 KiB',16:'16 MiB'}[b],f"{t['construction']['median_us']:,.2f}",f"{t['first_construction_us']:,.2f}",f"{pt['build_seconds_single_run']*1000:,.3f}"])
table(['Leaf bits','Patterns','C table','C median us','C first us','Python ms'],rows,[.65,.85,.95,1.05,1.05,1.47])
p('C entries contain full complex128 spectra, 16 bytes per complex output. Python stores object tuples and complex objects: about 3.3 KiB, 92 KiB and 43 MiB for 4, 8 and 16 bit leaves. C setup is allocation plus population, measured 11 times; the first construction is also shown. These are setup costs, not controlled cold-cache latency measurements.')
p('An 8-bit table needs only 256 entries. A 16-bit table needs 65,536 entries and adds random memory traffic; fewer arithmetic operations do not guarantee less time. To amortize a table, the number of transforms must exceed table setup time divided by the positive per-transform saving. If the warm variant loses, repeated use cannot amortize that loss.')

# Page 3

d.add_page_break();h('Recommended implementation')
for text in [
 'Start with a fixed 8-bit direct-index table, and also benchmark 4-bit and arithmetic codelet leaves. Pack the actual decimated FFT leaf into an integer key and index a flat read-only array. No hash table, canonical node IDs, LRU policy or per-node allocation is needed.',
 'Above the cutoff, run the ordinary FFT with no duplicate tracking. Reuse twiddles, input permutation and working buffers across calls. Select the best leaf method once per size and precision, rather than maintaining adaptive bookkeeping inside every FFT.',
 'For uniform random input, avoid caching large subproblems. At N = 65,536 the expected repeated fraction is about 96.875% for 8-bit subproblems, 3.06% for 16-bit subproblems and 0.000024% for 32-bit subproblems. A partial table of equally likely random patterns has no privileged frequent subset.',
 'An untested next candidate is a packed 8-bit real-input table. Store DC and Nyquist as real numbers plus three complex bins: eight doubles, 64 bytes per pattern, 16 KiB total. Reconstruct conjugate bins while writing the full output. Its extra reconstruction work needs measurement; no speedup is claimed here.'
]: d.add_paragraph(text,style='List Bullet')
h('Correctness and limits')
maxerr=max(c['correctness']['max_absolute_error'].values())
p(f"The C variants passed {first['correctness']['full_fft_cases']} complete FFT tests through N = 1,048,576, five random seeds per tested size and structured inputs, plus every 4, 8 and 16 bit table row. Worst complete-transform absolute error versus NumPy was {maxerr:.3g}. Python tests covered three seeds and sizes for all eight methods, exhaustive 4 and 8 bit rows, and structured inputs; worst error was 4.34e-13. ASan and UBSan checks passed; leak detection was unavailable under sandbox tracing.")
p('The algorithm is still O(N log N) for a fixed small cutoff. An 8-bit table eliminates three merge levels from online execution; it does not establish a new asymptotic FFT bound or a general-purpose speed advantage. Native code is a purpose-built binary-input microbenchmark, not a replacement claim for FFTW or other tuned libraries.')
p('SciPy and NumPy practical comparisons are retained in the JSON. Their calls include Python dispatch and output allocation, unlike the preallocated C kernel. SciPy rfft returns only N/2 + 1 bins, so its result is contextual and is not an equivalent full-output comparison. All use the same numerical precision.')
h('Reproduce the experiment')
p('The accompanying ZIP contains C source and an auto-compiling Python harness, the Python lookup experiment, the unchanged earlier prototype, and both native datasets. Requirements: Python 3, NumPy, SciPy and GCC on a POSIX platform. No installation or network access is performed by the scripts. Run commands are in README.md.')
p('Environment: AMD EPYC 9V74, Linux x86-64 cloud VM; GCC 14.2.0 with -O3 -march=native, double precision; NumPy 2.3.5 and SciPy 1.17.0. Results apply to this implementation and machine.')
p('Sources: the supplied fft_collision_research.py and fft_collision_research_notes.md; FFTW documentation on generated small codelets (https://fftw.org/fftw3_doc/Generating-your-own-code.html) and reusing plans (https://www.fftw.org/faq/section3.html). The binary lookup design and measurements here are an experiment, with no claim of novelty.')
for el in d.styles.element.xpath('.//w:pBdr') + d.element.xpath('.//w:pBdr'):
 el.getparent().remove(el)
d.save(out/'selective_fft_benchmark_report.docx')
print(out/'selective_fft_benchmark_report.docx')
